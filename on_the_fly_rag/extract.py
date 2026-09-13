"""Extract plain text from Office/PDF binaries for ingest.

Extractors are the only place binary formats are handled. Failures raise
``ExtractError`` so callers can log/skip without crashing the whole ingest.
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Callable, Dict, Optional, Set

logger = logging.getLogger(__name__)

# Extensions handled by extractors (not UTF-8 text reads).
BINARY_EXTENSIONS: Set[str] = {".pdf", ".docx", ".pptx"}


class ExtractError(Exception):
    """Raised when a binary document cannot be parsed into text."""

    def __init__(self, path: Path | str, reason: str) -> None:
        self.path = Path(path)
        self.reason = reason
        super().__init__(f"{self.path}: {reason}")


def extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover
        raise ExtractError(path, "pypdf is not installed") from e
    try:
        reader = PdfReader(str(path))
        parts: list[str] = []
        for i, page in enumerate(reader.pages):
            try:
                text = page.extract_text() or ""
            except Exception as page_err:  # noqa: BLE001
                logger.warning("PDF page %s extract failed in %s: %s", i, path, page_err)
                continue
            if text.strip():
                parts.append(text)
        joined = "\n\n".join(parts).strip()
        if not joined:
            raise ExtractError(path, "no extractable text in PDF")
        return joined
    except ExtractError:
        raise
    except Exception as e:  # noqa: BLE001
        raise ExtractError(path, f"PDF parse failed: {e}") from e


def extract_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as e:  # pragma: no cover
        raise ExtractError(path, "python-docx is not installed") from e
    try:
        doc = Document(str(path))
        parts: list[str] = []
        for para in doc.paragraphs:
            t = (para.text or "").strip()
            if t:
                parts.append(t)
        # Tables (optional, keep simple)
        for table in doc.tables:
            for row in table.rows:
                cells = [(c.text or "").strip() for c in row.cells]
                line = " | ".join(c for c in cells if c)
                if line:
                    parts.append(line)
        joined = "\n\n".join(parts).strip()
        if not joined:
            raise ExtractError(path, "no extractable text in DOCX")
        return joined
    except ExtractError:
        raise
    except Exception as e:  # noqa: BLE001
        raise ExtractError(path, f"DOCX parse failed: {e}") from e


def extract_pptx(path: Path) -> str:
    try:
        from pptx import Presentation
    except ImportError as e:  # pragma: no cover
        raise ExtractError(path, "python-pptx is not installed") from e
    try:
        prs = Presentation(str(path))
        parts: list[str] = []
        for idx, slide in enumerate(prs.slides, 1):
            slide_bits: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    t = (shape.text or "").strip()
                    if t:
                        slide_bits.append(t)
            notes = ""
            if slide.has_notes_slide:
                notes = (slide.notes_slide.notes_text_frame.text or "").strip()
            block = f"[Slide {idx}]\n" + "\n".join(slide_bits)
            if notes:
                block += f"\n[Notes]\n{notes}"
            if slide_bits or notes:
                parts.append(block)
        joined = "\n\n".join(parts).strip()
        if not joined:
            raise ExtractError(path, "no extractable text in PPTX")
        return joined
    except ExtractError:
        raise
    except Exception as e:  # noqa: BLE001
        raise ExtractError(path, f"PPTX parse failed: {e}") from e


_EXTRACTORS: Dict[str, Callable[[Path], str]] = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".pptx": extract_pptx,
}


def is_binary_document(path: Path | str) -> bool:
    return Path(path).suffix.lower() in BINARY_EXTENSIONS


def load_document(path: Path | str) -> str:
    """Load a file as plain text.

    Binary Office/PDF formats go through extractors; everything else is a
    UTF-8 text read (``errors=replace``). Raises ``ExtractError`` on failure.
    """
    path = Path(path)
    ext = path.suffix.lower()
    extractor = _EXTRACTORS.get(ext)
    if extractor is not None:
        return extractor(path)
    try:
        return path.read_text(encoding="utf-8", errors="replace")
    except OSError as e:
        raise ExtractError(path, f"text read failed: {e}") from e


def try_load_document(path: Path | str) -> Optional[str]:
    """Like ``load_document`` but logs and returns None on failure."""
    try:
        return load_document(path)
    except ExtractError as e:
        logger.warning("Skipping %s (%s)", e.path, e.reason)
        return None
